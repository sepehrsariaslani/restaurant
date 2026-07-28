"""POS feature pack endpoints.

This module implements the management/POS features that were missing from the
checklist, re-exported through ``restaurant.api`` (see the import at the bottom
of ``api.py``) so the frontend can reach them via
``/api/method/restaurant.api.<endpoint>``:

- bulk product operations (activate / deactivate / out-of-stock groups)
- Excel export & import of products
- barcode lookup for POS scanners
- combo (Product Bundle) management
- packaging fee settings + order charging
- work shift definitions + sales-by-shift report
- cash register settlement & closing (Restaurant Register Closing)
- extra BI reports (category / table / payment-method / product / shift sales)
- printer font settings
- per-user dashboard layout persistence
"""

import base64
import csv
import io
import json
import re
from collections import defaultdict
from html import escape as html_escape

import frappe
from frappe import _
from frappe.utils import cint, flt, get_datetime, getdate, now_datetime, today

from restaurant.api import (
	_bi_kpi,
	_compose_management_report,
	_default_uom,
	_ensure_management_access,
	_ensure_pos_shift_setting_fields,
	_get_currency,
	_get_default_item_price_rate,
	_get_default_selling_price_list_name,
	_get_single_setting,
	_has_column,
	_is_revenue_order,
	_json_safe_datetime,
	_management_collect_orders,
	_management_get_print_brand_settings,
	_management_resolve_item_name,
	_parse_json,
	_pos_shift_settings,
	_safe_div,
	_table_columns_from_rows,
)

REGISTER_CLOSING_DOCTYPE = "Restaurant Register Closing"
PRODUCT_EXPORT_COLUMNS = [
	"item_code",
	"item_name",
	"item_group",
	"category",
	"subcategory",
	"uom",
	"price",
	"barcode",
	"restaurant_enabled",
	"out_of_stock",
	"packaging_price",
	"short_desc",
	"description",
]
PRODUCT_BULK_ACTIONS = {"activate", "deactivate", "mark_out_of_stock", "mark_in_stock"}
PRINT_FONT_FAMILY_OPTIONS = ["Peyda", "Vazirmatn", "IRANSans", "Tahoma", "Sahel", "Shabnam", "Samim", "Arial"]
PRINT_RECEIPT_SCALE_OPTIONS = ["کوچک", "متوسط", "بزرگ"]
MANAGEMENT_FEATURE_PACK_REPORTS = {
	"category-sales",
	"table-sales",
	"payment-methods",
	"product-sales",
	"shift-sales",
}
PAYMENT_METHOD_LABELS = {
	"cash": "نقدی",
	"card": "کارتی / کارتخوان",
	"credit": "نسیه (اعتباری)",
	"other": "سایر",
	"unknown": "نامشخص",
}
DASHBOARD_WIDGET_KEYS = [
	"pos_shift",
	"kpis",
	"sales_financial",
	"time_trend",
	"cost_control",
	"performers",
	"channels",
	"crm",
	"menu",
	"lost_orders",
]

__all__ = [
	"MANAGEMENT_FEATURE_PACK_REPORTS",
	"fp_build_report_bi",
	# product ops
	"bulk_update_management_products",
	"get_management_product_by_barcode",
	"export_management_products_excel",
	"import_management_products_excel",
	# combos
	"list_management_combos",
	"get_management_combo_detail",
	"save_management_combo",
	"delete_management_combo",
	# packaging
	"get_management_packaging_settings",
	"set_management_packaging_settings",
	"fp_compute_order_packaging_fee",
	# shifts
	"get_management_work_shifts",
	"set_management_work_shifts",
	"fp_get_work_shifts",
	# register closing
	"get_management_register_closing_summary",
	"close_management_register",
	"list_management_register_closings",
	"get_management_register_closing_detail",
	# reports
	"get_management_report_category_sales",
	"get_management_report_table_sales",
	"get_management_report_payment_methods",
	"get_management_report_product_sales",
	"get_management_report_shift_sales",
	# print fonts
	"get_management_print_font_settings",
	"set_management_print_font_settings",
	# dashboard layout
	"get_management_dashboard_layout",
	"set_management_dashboard_layout",
]


# ---------------------------------------------------------------------------
# Custom field provisioning (runtime, idempotent)
# ---------------------------------------------------------------------------


def _fp_resolve_insert_after(doctype_name, candidates):
	"""Return the first existing fieldname from candidates to anchor new fields."""
	try:
		meta = frappe.get_meta(doctype_name)
	except Exception:
		return ""
	for candidate in candidates or []:
		try:
			if candidate and meta.has_field(candidate):
				return candidate
		except Exception:
			continue
	fields = [field.fieldname for field in (getattr(meta, "fields", None) or [])]
	return fields[-1] if fields else ""


def _fp_ensure_custom_fields(doctype_name, field_defs, anchor_candidates=None):
	"""Idempotently create/update Custom Field records for ``doctype_name``."""
	if not frappe.db.exists("DocType", doctype_name):
		return

	anchor = _fp_resolve_insert_after(doctype_name, anchor_candidates or [])
	previous_fieldname = anchor
	changed_any = False

	for field_def in field_defs:
		payload_def = dict(field_def)
		if previous_fieldname and "insert_after" not in payload_def:
			payload_def["insert_after"] = previous_fieldname
		fieldname = payload_def.get("fieldname")
		if not fieldname:
			continue

		existing_name = frappe.db.get_value(
			"Custom Field",
			{"dt": doctype_name, "fieldname": fieldname},
			"name",
		)
		payload = {
			"doctype": "Custom Field",
			"dt": doctype_name,
			"module": "Restaurant",
			**payload_def,
		}

		if existing_name:
			doc = frappe.get_doc("Custom Field", existing_name)
			changed = False
			for key, value in payload.items():
				if key in {"doctype", "insert_after"}:
					continue
				if doc.get(key) != value:
					doc.set(key, value)
					changed = True
			if changed:
				doc.save(ignore_permissions=True)
				changed_any = True
		else:
			frappe.get_doc(payload).insert(ignore_permissions=True)
			changed_any = True

		previous_fieldname = fieldname

	if changed_any:
		frappe.clear_cache(doctype=doctype_name)


def _fp_ensure_item_ops_fields():
	"""Custom fields on Item: out-of-stock flag + per-item packaging price."""
	_fp_ensure_custom_fields(
		"Item",
		[
			{
				"fieldname": "restaurant_out_of_stock",
				"label": "ناموجود (اتمام موقت)",
				"fieldtype": "Check",
				"default": "0",
				"in_standard_filter": 1,
			},
			{
				"fieldname": "restaurant_packaging_price",
				"label": "هزینه بسته بندی",
				"fieldtype": "Currency",
				"default": "0",
			},
		],
		anchor_candidates=("restaurant_enabled", "disabled"),
	)


def _fp_ensure_ops_setting_fields():
	"""Custom fields on Restaurant Web Settings: packaging, work shifts."""
	# Make sure the shift-settings section (owned by api.py) already exists so
	# our section can anchor after it on existing sites.
	try:
		_ensure_pos_shift_setting_fields()
	except Exception:
		pass

	_fp_ensure_custom_fields(
		"Restaurant Web Settings",
		[
			{
				"fieldname": "restaurant_ops_feature_section",
				"label": "تنظیمات بسته بندی، شیفت و عملیات صندوق",
				"fieldtype": "Section Break",
			},
			{
				"fieldname": "restaurant_packaging_enabled",
				"label": "فعال سازی هزینه بسته بندی",
				"fieldtype": "Check",
				"default": "0",
			},
			{
				"fieldname": "restaurant_packaging_flat_fee",
				"label": "هزینه ثابت بسته بندی هر سفارش",
				"fieldtype": "Currency",
				"default": "0",
			},
			{
				"fieldname": "restaurant_packaging_per_item",
				"label": "محاسبه هزینه بسته بندی به ازای هر محصول",
				"fieldtype": "Check",
				"default": "1",
			},
			{
				"fieldname": "restaurant_packaging_takeaway",
				"label": "اعمال روی سفارش های بیرون بر",
				"fieldtype": "Check",
				"default": "1",
			},
			{
				"fieldname": "restaurant_packaging_delivery",
				"label": "اعمال روی سفارش های ارسالی",
				"fieldtype": "Check",
				"default": "1",
			},
			{
				"fieldname": "restaurant_packaging_dine_in",
				"label": "اعمال روی سفارش های حضوری (سالن)",
				"fieldtype": "Check",
				"default": "0",
			},
			{
				"fieldname": "restaurant_packaging_label",
				"label": "عنوان هزینه بسته بندی در فاکتور",
				"fieldtype": "Data",
				"default": "هزینه بسته بندی",
			},
			{
				"fieldname": "restaurant_work_shifts_json",
				"label": "تعریف شیفت های کاری (JSON)",
				"fieldtype": "Long Text",
			},
		],
		anchor_candidates=("restaurant_pos_closing_note_template", "primary_cta_label"),
	)


def _fp_ensure_sales_order_packaging_field():
	_fp_ensure_custom_fields(
		"Sales Order",
		[
			{
				"fieldname": "restaurant_packaging_fee",
				"label": "هزینه بسته بندی",
				"fieldtype": "Currency",
				"default": "0",
			},
		],
		anchor_candidates=("restaurant_delivery_fee", "restaurant_note", "restaurant_status"),
	)


def _fp_ensure_ops_ready():
	_fp_ensure_item_ops_fields()
	_fp_ensure_ops_setting_fields()
	_fp_ensure_sales_order_packaging_field()


# ---------------------------------------------------------------------------
# Bulk product operations
# ---------------------------------------------------------------------------


def _fp_normalize_item_names(item_names):
	if isinstance(item_names, str):
		parsed = _parse_json(item_names, None)
		if isinstance(parsed, list):
			item_names = parsed
		else:
			item_names = [chunk.strip() for chunk in item_names.split(",") if chunk.strip()]
	names = []
	for raw in item_names or []:
		name = (raw or "").strip()
		if name and name not in names:
			names.append(name)
	return names[:500]


@frappe.whitelist()
def bulk_update_management_products(item_names=None, action=None):
	"""Bulk activate/deactivate or mark products in/out of stock."""
	_ensure_management_access()
	action = (action or "").strip()
	if action not in PRODUCT_BULK_ACTIONS:
		frappe.throw(
			_("Invalid bulk action: {0}. Allowed: {1}").format(
				action or "-", ", ".join(sorted(PRODUCT_BULK_ACTIONS))
			)
		)

	names = _fp_normalize_item_names(item_names)
	if not names:
		frappe.throw(_("No products selected for the bulk operation."))

	_fp_ensure_item_ops_fields()

	results = []
	updated = 0
	for name in names:
		try:
			target = _management_resolve_item_name(name)
			if action == "activate":
				updates = {"restaurant_enabled": 1}
			elif action == "deactivate":
				updates = {"restaurant_enabled": 0}
			elif action == "mark_out_of_stock":
				updates = {"restaurant_out_of_stock": 1}
			else:
				updates = {"restaurant_out_of_stock": 0}

			applied = {key: value for key, value in updates.items() if _has_column("Item", key)}
			if not applied:
				frappe.throw(_("Required custom fields are missing on Item."))

			frappe.db.set_value("Item", target, applied, update_modified=True)
			updated += 1
			results.append({"item_name": target, "status": "ok", **applied})
		except Exception as error:
			results.append({"item_name": name, "status": "error", "message": str(error)})

	frappe.db.commit()
	return {
		"status": "success",
		"action": action,
		"updated": updated,
		"failed": len(results) - updated,
		"results": results,
	}


# ---------------------------------------------------------------------------
# Barcode lookup
# ---------------------------------------------------------------------------


def _fp_first_barcode_map(item_names):
	"""Return {item_name: first barcode} for the given items."""
	names = [name for name in (item_names or []) if name]
	if not names or not frappe.db.exists("DocType", "Item Barcode"):
		return {}
	rows = frappe.get_all(
		"Item Barcode",
		fields=["parent", "barcode"],
		filters={"parent": ["in", names]},
		order_by="idx asc",
		ignore_permissions=True,
		limit_page_length=len(names) * 3,
	)
	barcode_map = {}
	for row in rows:
		parent = row.get("parent")
		barcode = (row.get("barcode") or "").strip()
		if parent and barcode and parent not in barcode_map:
			barcode_map[parent] = barcode
	return barcode_map


@frappe.whitelist()
def get_management_product_by_barcode(barcode=None):
	"""Resolve a product for POS barcode scanners."""
	_ensure_management_access()
	code = (barcode or "").strip()
	if not code:
		frappe.throw(_("Barcode is required."))

	item_name = ""
	if frappe.db.exists("DocType", "Item Barcode"):
		item_name = frappe.db.get_value("Item Barcode", {"barcode": code}, "parent") or ""
	if not item_name:
		item_name = frappe.db.get_value("Item", {"item_code": code}, "name") or ""
	if not item_name and frappe.db.exists("Item", code):
		item_name = code

	if not item_name:
		return {"status": "not_found", "barcode": code, "item": None}

	fields = ["name", "item_code", "item_name", "item_group", "stock_uom", "disabled", "standard_rate"]
	for optional in ("restaurant_slug", "restaurant_enabled", "restaurant_out_of_stock"):
		if _has_column("Item", optional):
			fields.append(optional)
	row = frappe.get_value("Item", item_name, fields, as_dict=True) or {}
	price = _get_default_item_price_rate({"name": row.get("name"), "item_code": row.get("item_code")})
	if price is None:
		price = flt(row.get("standard_rate") or 0)

	return {
		"status": "success",
		"barcode": code,
		"item": {
			"name": row.get("name") or "",
			"item_code": row.get("item_code") or row.get("name") or "",
			"item_name": row.get("item_name") or row.get("name") or "",
			"title": row.get("item_name") or row.get("name") or "",
			"slug": row.get("restaurant_slug") or "",
			"item_group": row.get("item_group") or "",
			"stock_uom": row.get("stock_uom") or "",
			"disabled": cint(row.get("disabled") or 0),
			"is_active": cint(row.get("restaurant_enabled") or 0) if "restaurant_enabled" in row else 1,
			"out_of_stock": cint(row.get("restaurant_out_of_stock") or 0),
			"base_price": flt(price),
		},
	}


# ---------------------------------------------------------------------------
# Excel export / import
# ---------------------------------------------------------------------------


def _fp_product_export_rows(category=None, include_disabled=0):
	include_disabled = cint(include_disabled)
	filters = {}
	if _has_column("Item", "variant_of"):
		filters["variant_of"] = ["in", ["", None]]
	if not include_disabled:
		filters["disabled"] = 0
	if category:
		filters["restaurant_category" if _has_column("Item", "restaurant_category") else "item_group"] = category

	fields = ["name", "item_code", "item_name", "item_group", "stock_uom", "disabled", "description"]
	for optional in (
		"restaurant_category",
		"restaurant_subcategory",
		"restaurant_enabled",
		"restaurant_out_of_stock",
		"restaurant_packaging_price",
		"restaurant_short_desc",
	):
		if _has_column("Item", optional):
			fields.append(optional)

	rows = frappe.get_all(
		"Item",
		filters=filters,
		fields=fields,
		order_by="item_name asc",
		ignore_permissions=True,
		limit_page_length=20000,
	)

	names = [row.get("name") for row in rows if row.get("name")]
	barcode_map = _fp_first_barcode_map(names)

	price_map = {}
	price_list = (_get_default_selling_price_list_name() or "").strip()
	if names and price_list and frappe.db.exists("DocType", "Item Price"):
		price_rows = frappe.get_all(
			"Item Price",
			fields=["item_code", "price_list_rate"],
			filters={"item_code": ["in", names], "price_list": price_list, "selling": 1},
			ignore_permissions=True,
			limit_page_length=20000,
		)
		for price_row in price_rows:
			price_map[price_row.get("item_code")] = flt(price_row.get("price_list_rate"))

	data = [list(PRODUCT_EXPORT_COLUMNS)]
	for row in rows:
		data.append(
			[
				row.get("item_code") or row.get("name") or "",
				row.get("item_name") or "",
				row.get("item_group") or "",
				row.get("restaurant_category") or "",
				row.get("restaurant_subcategory") or "",
				row.get("stock_uom") or "",
				price_map.get(row.get("name"), flt(row.get("standard_rate") or 0)),
				barcode_map.get(row.get("name"), ""),
				cint(row.get("restaurant_enabled") or 0) if "restaurant_enabled" in row else "",
				cint(row.get("restaurant_out_of_stock") or 0),
				flt(row.get("restaurant_packaging_price") or 0),
				row.get("restaurant_short_desc") or "",
				row.get("description") or "",
			]
		)
	return data


@frappe.whitelist()
def export_management_products_excel(category=None, include_disabled=0):
	"""Export the product catalog as an .xlsx file and return its URL."""
	_ensure_management_access()
	_fp_ensure_item_ops_fields()

	data = _fp_product_export_rows(category=category, include_disabled=include_disabled)

	try:
		from frappe.utils.xlsxutils import make_xlsx
	except Exception:
		frappe.throw(_("Excel export is not available on this server (frappe xlsxutils missing)."))

	xlsx_file = make_xlsx(data, "Products")
	file_name = "restaurant-products-{}.xlsx".format(now_datetime().strftime("%Y%m%d-%H%M%S"))
	file_doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": file_name,
			"content": xlsx_file.getvalue(),
			"is_private": 1,
			"folder": "Home",
		}
	).insert(ignore_permissions=True)
	frappe.db.commit()

	return {
		"status": "success",
		"file_url": file_doc.file_url,
		"file_name": file_doc.file_name,
		"rows": max(len(data) - 1, 0),
		"columns": list(PRODUCT_EXPORT_COLUMNS),
	}


def _fp_load_upload_content(file_url=None, content_base64=None):
	if content_base64:
		try:
			return base64.b64decode(content_base64)
		except Exception:
			frappe.throw(_("Invalid base64 file content."))

	file_url = (file_url or "").strip()
	if not file_url:
		frappe.throw(_("No file was provided. Upload a file or pass content_base64."))

	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		frappe.throw(_("Uploaded file not found: {0}").format(file_url))
	file_doc = frappe.get_doc("File", file_name)
	content = file_doc.get_content()
	if isinstance(content, str):
		return content.encode("utf-8")
	return content


def _fp_parse_tabular_rows(raw, file_name=""):
	name = (file_name or "").lower()
	if name.endswith(".csv"):
		text = raw.decode("utf-8-sig", "ignore")
		return [[("" if cell is None else str(cell).strip()) for cell in row] for row in csv.reader(io.StringIO(text))]

	try:
		from openpyxl import load_workbook
	except ImportError:
		frappe.throw(
			_("Excel import requires openpyxl on the server. Please upload a CSV file instead.")
		)
	workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
	sheet = workbook.active
	return [[("" if cell is None else str(cell).strip()) for cell in row] for row in sheet.iter_rows(values_only=True)]


PRODUCT_IMPORT_HEADER_ALIASES = {
	"item_code": {"item_code", "code", "کد", "کد کالا"},
	"item_name": {"item_name", "name", "title", "نام", "نام کالا"},
	"item_group": {"item_group", "group", "گروه", "گروه کالا"},
	"category": {"category", "restaurant_category", "دسته", "دسته بندی"},
	"subcategory": {"subcategory", "restaurant_subcategory", "زیردسته"},
	"uom": {"uom", "stock_uom", "واحد"},
	"price": {"price", "rate", "قیمت"},
	"barcode": {"barcode", "بارکد"},
	"restaurant_enabled": {"restaurant_enabled", "is_active", "active", "enabled", "فعال"},
	"out_of_stock": {"out_of_stock", "ناموجود", "اتمام", "اتمام موجودی"},
	"packaging_price": {"packaging_price", "هزینه بسته بندی", "بسته بندی"},
	"short_desc": {"short_desc", "توضیح کوتاه"},
	"description": {"description", "توضیحات"},
}


def _fp_normalize_import_headers(header_row):
	mapping = {}
	for alias_key, aliases in PRODUCT_IMPORT_HEADER_ALIASES.items():
		normalized_aliases = {alias.strip().lower() for alias in aliases}
		for idx, cell in enumerate(header_row or []):
			if (cell or "").strip().lower() in normalized_aliases:
				mapping[idx] = alias_key
	return mapping


def _fp_parse_import_flag(value):
	text = (value or "").strip().lower()
	if not text:
		return None
	if text in {"1", "true", "yes", "y", "فعال", "بله", "✓", "✔"}:
		return 1
	if text in {"0", "false", "no", "n", "غیرفعال", "خیر", "✗", "✘"}:
		return 0
	parsed = _parse_json(text, None)
	if isinstance(parsed, (int, float)):
		return 1 if parsed else 0
	return None


def _fp_resolve_item_group(raw_value):
	value = (raw_value or "").strip()
	if value and frappe.db.exists("Item Group", value):
		return value
	if value:
		by_title = frappe.db.get_value("Item Group", {"item_group_name": value}, "name")
		if by_title:
			return by_title
	for candidate in ("Products", "All Item Groups"):
		if frappe.db.exists("Item Group", candidate):
			return candidate
	fallback = frappe.get_all("Item Group", filters={"is_group": 0}, pluck="name", limit_page_length=1)
	return fallback[0] if fallback else "All Item Groups"


def _fp_upsert_item_price(item_code, rate):
	rate = flt(rate)
	if rate <= 0 or not item_code or not frappe.db.exists("DocType", "Item Price"):
		return
	price_list = (_get_default_selling_price_list_name() or "").strip()
	if not price_list:
		return
	existing = frappe.db.get_value(
		"Item Price",
		{"item_code": item_code, "price_list": price_list, "selling": 1},
		"name",
	)
	if existing:
		frappe.db.set_value("Item Price", existing, "price_list_rate", rate, update_modified=True)
		return
	frappe.get_doc(
		{
			"doctype": "Item Price",
			"item_code": item_code,
			"price_list": price_list,
			"selling": 1,
			"buying": 0,
			"price_list_rate": rate,
			"currency": _get_currency(),
			"uom": frappe.db.get_value("Item", item_code, "stock_uom"),
		}
	).insert(ignore_permissions=True)


def _fp_upsert_item_barcode(item_name, barcode):
	barcode = (barcode or "").strip()
	if not barcode or not frappe.db.exists("DocType", "Item Barcode"):
		return
	existing = frappe.db.get_value("Item Barcode", {"parent": item_name, "barcode": barcode}, "name")
	if existing:
		return
	doc = frappe.get_doc("Item", item_name)
	doc.append("barcodes", {"barcode": barcode, "uom": doc.get("stock_uom")})
	doc.save(ignore_permissions=True)


def _fp_import_product_record(record, update_existing=1, dry_run=0):
	update_existing = cint(update_existing)
	dry_run = cint(dry_run)

	item_code = (record.get("item_code") or "").strip()
	if not item_code:
		return {"status": "error", "message": _("Item code is missing.")}

	item_name = (record.get("item_name") or "").strip() or item_code
	item_group = _fp_resolve_item_group(record.get("item_group") or record.get("category") or "")
	uom = (record.get("uom") or "").strip() or _default_uom()
	price = flt(record.get("price") or 0)
	barcode = (record.get("barcode") or "").strip()
	enabled_flag = _fp_parse_import_flag(record.get("restaurant_enabled"))
	out_of_stock_flag = _fp_parse_import_flag(record.get("out_of_stock"))
	packaging_price = flt(record.get("packaging_price") or 0)
	description = (record.get("description") or "").strip()
	short_desc = (record.get("short_desc") or "").strip()

	existing_name = frappe.db.get_value("Item", {"item_code": item_code}, "name")
	if not existing_name and frappe.db.exists("Item", item_code):
		existing_name = item_code

	if dry_run:
		return {"status": "ok", "mode": "update" if existing_name else "create", "item_code": item_code}

	if existing_name and not update_existing:
		return {"status": "skipped", "message": _("Item already exists and update_existing is off.")}

	if existing_name:
		doc = frappe.get_doc("Item", existing_name)
		doc.set("item_name", item_name)
		if item_group:
			doc.set("item_group", item_group)
	else:
		doc = frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": item_code,
				"item_name": item_name,
				"item_group": item_group,
				"stock_uom": uom,
				"is_sales_item": 1,
				"is_stock_item": 0,
				"disabled": 0,
			}
		)

	guarded_values = {}
	if enabled_flag is not None and _has_column("Item", "restaurant_enabled"):
		guarded_values["restaurant_enabled"] = enabled_flag
	elif not existing_name and _has_column("Item", "restaurant_enabled"):
		guarded_values["restaurant_enabled"] = 1
	if out_of_stock_flag is not None and _has_column("Item", "restaurant_out_of_stock"):
		guarded_values["restaurant_out_of_stock"] = out_of_stock_flag
	if packaging_price and _has_column("Item", "restaurant_packaging_price"):
		guarded_values["restaurant_packaging_price"] = packaging_price
	if short_desc and _has_column("Item", "restaurant_short_desc"):
		guarded_values["restaurant_short_desc"] = short_desc
	if description:
		guarded_values["description"] = description
	category_value = (record.get("category") or "").strip()
	if category_value and _has_column("Item", "restaurant_category"):
		category_group = _fp_resolve_item_group(category_value)
		if category_group:
			guarded_values["restaurant_category"] = category_group

	for key, value in guarded_values.items():
		doc.set(key, value)

	if existing_name:
		doc.save(ignore_permissions=True)
	else:
		doc.insert(ignore_permissions=True)

	_fp_upsert_item_price(doc.name, price)
	_fp_upsert_item_barcode(doc.name, barcode)

	return {"status": "ok", "mode": "update" if existing_name else "create", "item_code": doc.name}


@frappe.whitelist()
def import_management_products_excel(
	file_url=None, content_base64=None, file_name=None, update_existing=1, dry_run=0
):
	"""Bulk import products from an uploaded .xlsx/.csv file."""
	_ensure_management_access()
	_fp_ensure_item_ops_fields()

	raw = _fp_load_upload_content(file_url=file_url, content_base64=content_base64)
	fallback_name = file_name or (file_url or "")
	rows = _fp_parse_tabular_rows(raw, file_name=fallback_name)
	if len(rows) < 2:
		frappe.throw(_("The uploaded file has no data rows."))

	header_mapping = _fp_normalize_import_headers(rows[0])
	if not header_mapping:
		frappe.throw(
			_("Could not detect the column headers. Expected columns like: {0}").format(
				", ".join(PRODUCT_EXPORT_COLUMNS)
			)
		)

	records = []
	for row in rows[1:]:
		if not any((cell or "").strip() for cell in row):
			continue
		record = {}
		for idx, key in header_mapping.items():
			if idx < len(row):
				record[key] = row[idx]
		records.append(record)

	if not records:
		frappe.throw(_("The uploaded file has no data rows."))

	summary = {"created": 0, "updated": 0, "skipped": 0, "errors": []}
	for index, record in enumerate(records, start=2):
		try:
			result = _fp_import_product_record(record, update_existing=update_existing, dry_run=dry_run)
			if result.get("status") == "skipped":
				summary["skipped"] += 1
			elif result.get("mode") == "update":
				summary["updated"] += 1
			else:
				summary["created"] += 1
		except Exception as error:
			summary["errors"].append({"row": index, "item_code": record.get("item_code") or "", "message": str(error)})

	if not cint(dry_run):
		frappe.db.commit()

	return {
		"status": "success",
		"dry_run": cint(dry_run),
		"total_rows": len(records),
		**summary,
	}


# ---------------------------------------------------------------------------
# Combo management (ERPNext Product Bundle wrapper)
# ---------------------------------------------------------------------------


def _fp_combo_bundle_available():
	return frappe.db.exists("DocType", "Product Bundle") and frappe.db.exists(
		"DocType", "Product Bundle Item"
	)


def _fp_serialize_combo(bundle_doc):
	parent_item = frappe.db.get_value(
		"Item",
		bundle_doc.get("new_item_code"),
		["name", "item_name", "disabled", "standard_rate"],
		as_dict=True,
	)
	price = None
	if parent_item:
		price = _get_default_item_price_rate({"name": parent_item.get("name"), "item_code": parent_item.get("name")})
	items = []
	for row in bundle_doc.get("items") or []:
		component = frappe.db.get_value("Item", row.get("item_code"), ["item_name"], as_dict=True) or {}
		items.append(
			{
				"item_code": row.get("item_code") or "",
				"item_name": component.get("item_name") or row.get("item_code") or "",
				"qty": flt(row.get("qty") or 0),
				"uom": row.get("uom") or "",
				"description": row.get("description") or "",
			}
		)
	return {
		"name": bundle_doc.get("name"),
		"combo_item": bundle_doc.get("new_item_code") or bundle_doc.get("name"),
		"item_name": (parent_item or {}).get("item_name") or bundle_doc.get("new_item_code"),
		"disabled": cint((parent_item or {}).get("disabled") or 0),
		"price": flt(price) if price is not None else flt((parent_item or {}).get("standard_rate") or 0),
		"items": items,
	}


def _fp_assert_combo_feature():
	if not _fp_combo_bundle_available():
		frappe.throw(
			_("Product Bundle doctype is not available. Make sure ERPNext stock module is installed.")
		)


@frappe.whitelist()
def list_management_combos(search=None):
	_ensure_management_access()
	_fp_assert_combo_feature()

	filters = {}
	search = (search or "").strip()
	if search:
		filters["new_item_code"] = ["like", f"%{search}%"]

	names = frappe.get_all(
		"Product Bundle",
		filters=filters,
		pluck="name",
		order_by="modified desc",
		limit_page_length=200,
		ignore_permissions=True,
	)
	combos = []
	for name in names:
		doc = frappe.get_doc("Product Bundle", name)
		combos.append(_fp_serialize_combo(doc))
	return {"combos": combos}


@frappe.whitelist()
def get_management_combo_detail(combo_item=None):
	_ensure_management_access()
	_fp_assert_combo_feature()
	combo_item = (combo_item or "").strip()
	if not combo_item or not frappe.db.exists("Product Bundle", combo_item):
		frappe.throw(_("Combo (Product Bundle) not found."), frappe.DoesNotExistError)
	doc = frappe.get_doc("Product Bundle", combo_item)
	return {"combo": _fp_serialize_combo(doc)}


@frappe.whitelist()
def save_management_combo(payload=None):
	"""Create or update a combo: a sellable parent item plus child products."""
	_ensure_management_access()
	_fp_assert_combo_feature()

	data = _parse_json(payload, {})
	if not isinstance(data, dict):
		frappe.throw(_("Invalid payload format."))

	combo_item = (data.get("combo_item") or data.get("item_code") or "").strip()
	if not combo_item:
		frappe.throw(_("Combo item (parent product) is required."))
	if not frappe.db.exists("Item", combo_item):
		frappe.throw(_("Combo parent item not found: {0}").format(combo_item))

	components = data.get("items") or []
	if not isinstance(components, list) or not components:
		frappe.throw(_("A combo needs at least one component product."))

	normalized_items = []
	for component in components:
		component_code = (component.get("item_code") or component.get("item") or "").strip()
		if not component_code:
			continue
		if component_code == combo_item:
			frappe.throw(_("A combo cannot contain itself."))
		if not frappe.db.exists("Item", component_code):
			frappe.throw(_("Component item not found: {0}").format(component_code))
		qty = flt(component.get("qty") or 0)
		if qty <= 0:
			qty = 1
		normalized_items.append(
			{
				"item_code": component_code,
				"qty": qty,
				"uom": (component.get("uom") or "").strip()
				or frappe.db.get_value("Item", component_code, "stock_uom"),
				"description": (component.get("description") or "").strip(),
			}
		)
	if not normalized_items:
		frappe.throw(_("A combo needs at least one valid component product."))
	if frappe.db.exists("Product Bundle", {"new_item_code": ["in", [row["item_code"] for row in normalized_items]]}):
		frappe.throw(_("A component cannot itself be a combo parent."))

	if frappe.db.exists("Product Bundle", combo_item):
		doc = frappe.get_doc("Product Bundle", combo_item)
		doc.set("items", [])
	else:
		doc = frappe.get_doc({"doctype": "Product Bundle", "new_item_code": combo_item, "items": []})

	for component in normalized_items:
		doc.append("items", component)
	doc.save(ignore_permissions=True)
	frappe.db.commit()

	return {"status": "success", "combo": _fp_serialize_combo(doc)}


@frappe.whitelist()
def delete_management_combo(combo_item=None):
	_ensure_management_access()
	_fp_assert_combo_feature()
	combo_item = (combo_item or "").strip()
	if not combo_item or not frappe.db.exists("Product Bundle", combo_item):
		frappe.throw(_("Combo (Product Bundle) not found."), frappe.DoesNotExistError)
	frappe.delete_doc("Product Bundle", combo_item, ignore_permissions=True)
	frappe.db.commit()
	return {"status": "success", "combo_item": combo_item}


# ---------------------------------------------------------------------------
# Packaging fee settings + computation
# ---------------------------------------------------------------------------


def _fp_packaging_apply_modes():
	modes = []
	for mode, fieldname in (
		("takeaway", "restaurant_packaging_takeaway"),
		("delivery", "restaurant_packaging_delivery"),
		("dine_in", "restaurant_packaging_dine_in"),
	):
		if cint(_get_single_setting("Restaurant Web Settings", fieldname, 0)) == 1:
			modes.append(mode)
	return modes


def _fp_packaging_settings_payload():
	return {
		"enabled": cint(_get_single_setting("Restaurant Web Settings", "restaurant_packaging_enabled", 0)) == 1,
		"flat_fee": flt(_get_single_setting("Restaurant Web Settings", "restaurant_packaging_flat_fee", 0)),
		"per_item": cint(_get_single_setting("Restaurant Web Settings", "restaurant_packaging_per_item", 1)) == 1,
		"apply_modes": _fp_packaging_apply_modes(),
		"label": (
			_get_single_setting("Restaurant Web Settings", "restaurant_packaging_label", "") or _("Packaging Fee")
		),
		"currency": _get_currency(),
	}


def fp_compute_order_packaging_fee(order_type, item_qty_map=None):
	"""Server-side packaging fee for an order (flat fee + per-item prices)."""
	settings = _fp_packaging_settings_payload()
	if not settings.get("enabled"):
		return 0.0

	mode = (order_type or "").strip().lower()
	if mode not in (settings.get("apply_modes") or []):
		return 0.0

	fee = flt(settings.get("flat_fee") or 0)
	if settings.get("per_item") and item_qty_map and _has_column("Item", "restaurant_packaging_price"):
		item_names = [name for name in item_qty_map.keys() if name]
		if item_names:
			rows = frappe.get_all(
				"Item",
				filters={"name": ["in", item_names]},
				fields=["name", "restaurant_packaging_price"],
				ignore_permissions=True,
				limit_page_length=len(item_names) + 5,
			)
			price_map = {row.get("name"): flt(row.get("restaurant_packaging_price") or 0) for row in rows}
			for item_name, qty in item_qty_map.items():
				fee += price_map.get(item_name, 0.0) * flt(qty)

	return flt(fee)


@frappe.whitelist()
def get_management_packaging_settings():
	_ensure_management_access()
	_fp_ensure_ops_setting_fields()
	return _fp_packaging_settings_payload()


@frappe.whitelist()
def set_management_packaging_settings(payload=None):
	_ensure_management_access()
	_fp_ensure_ops_setting_fields()
	data = _parse_json(payload, {})
	if not isinstance(data, dict):
		frappe.throw(_("Invalid payload format."))

	def _set(fieldname, value):
		frappe.db.set_single_value("Restaurant Web Settings", fieldname, value, update_modified=True)

	if "enabled" in data:
		_set("restaurant_packaging_enabled", 1 if cint(data.get("enabled")) else 0)
	if "flat_fee" in data:
		_set("restaurant_packaging_flat_fee", max(flt(data.get("flat_fee") or 0), 0))
	if "per_item" in data:
		_set("restaurant_packaging_per_item", 1 if cint(data.get("per_item")) else 0)
	if "apply_modes" in data and isinstance(data.get("apply_modes"), (list, tuple)):
		modes = {str(mode).strip().lower() for mode in data.get("apply_modes")}
		_set("restaurant_packaging_takeaway", 1 if "takeaway" in modes else 0)
		_set("restaurant_packaging_delivery", 1 if "delivery" in modes else 0)
		_set("restaurant_packaging_dine_in", 1 if "dine_in" in modes else 0)
	for mode in ("takeaway", "delivery", "dine_in"):
		key = f"apply_{mode}"
		if key in data:
			_set(f"restaurant_packaging_{mode}", 1 if cint(data.get(key)) else 0)
	if "label" in data:
		label = (data.get("label") or "").strip() or _("Packaging Fee")
		_set("restaurant_packaging_label", label)

	frappe.db.commit()
	return {"status": "success", "settings": _fp_packaging_settings_payload()}


# ---------------------------------------------------------------------------
# Work shifts
# ---------------------------------------------------------------------------


def _fp_default_work_shifts():
	return [
		{"key": "morning", "label": "صبح", "start": "06:00", "end": "12:00"},
		{"key": "noon", "label": "ظهر", "start": "12:00", "end": "18:00"},
		{"key": "night", "label": "شب", "start": "18:00", "end": "24:00"},
	]


def _fp_normalize_shift_time(value):
	text = str(value or "").strip()
	match = re.match(r"^(\d{1,2}):(\d{2})(?::\d{2})?$", text)
	if not match:
		return ""
	hours = cint(match.group(1))
	minutes = cint(match.group(2))
	if hours > 24 or minutes > 59 or (hours == 24 and minutes > 0):
		return ""
	return f"{hours:02d}:{minutes:02d}"


def _fp_normalize_work_shifts(raw_shifts):
	normalized = []
	for idx, shift in enumerate(raw_shifts or []):
		if not isinstance(shift, dict):
			continue
		start = _fp_normalize_shift_time(shift.get("start"))
		end = _fp_normalize_shift_time(shift.get("end"))
		if not start or not end:
			continue
		key = (shift.get("key") or "").strip() or f"shift_{idx + 1}"
		label = (shift.get("label") or "").strip() or key
		normalized.append({"key": key, "label": label, "start": start, "end": end})
	return normalized


def fp_get_work_shifts():
	"""Return configured work shifts (with sane defaults)."""
	try:
		raw = _get_single_setting("Restaurant Web Settings", "restaurant_work_shifts_json", "")
	except Exception:
		raw = ""
	parsed = _parse_json(raw or "[]", [])
	shifts = _fp_normalize_work_shifts(parsed if isinstance(parsed, list) else [])
	return shifts or _fp_default_work_shifts()


def _fp_shift_minutes(value):
	normalized = _fp_normalize_shift_time(value)
	if not normalized:
		return None
	hours, minutes = normalized.split(":")
	return cint(hours) * 60 + cint(minutes)


def _fp_shift_for_datetime(shifts, dt):
	try:
		dt = get_datetime(dt)
	except Exception:
		return None
	minutes = dt.hour * 60 + dt.minute
	for shift in shifts:
		start = _fp_shift_minutes(shift.get("start"))
		end = _fp_shift_minutes(shift.get("end"))
		if start is None or end is None:
			continue
		if start < end and start <= minutes < end:
			return shift
		if start >= end and (minutes >= start or minutes < end):
			return shift
	return None


@frappe.whitelist()
def get_management_work_shifts():
	_ensure_management_access()
	_fp_ensure_ops_setting_fields()
	return {"shifts": fp_get_work_shifts(), "defaults": _fp_default_work_shifts()}


@frappe.whitelist()
def set_management_work_shifts(payload=None):
	_ensure_management_access()
	_fp_ensure_ops_setting_fields()
	data = _parse_json(payload, {})
	shifts_raw = data.get("shifts") if isinstance(data, dict) else data
	shifts = _fp_normalize_work_shifts(shifts_raw if isinstance(shifts_raw, list) else [])
	if not shifts:
		frappe.throw(_("At least one valid shift (with start/end times) is required."))
	if len(shifts) > 8:
		frappe.throw(_("A maximum of 8 work shifts can be defined."))

	frappe.db.set_single_value(
		"Restaurant Web Settings",
		"restaurant_work_shifts_json",
		json.dumps(shifts, ensure_ascii=False),
		update_modified=True,
	)
	frappe.db.commit()
	return {"status": "success", "shifts": shifts}


# ---------------------------------------------------------------------------
# Cash register settlement & closing
# ---------------------------------------------------------------------------


def _fp_register_closing_available():
	return frappe.db.exists("DocType", REGISTER_CLOSING_DOCTYPE)


def _fp_last_register_closing(cashier=None, pos_profile=None):
	if not _fp_register_closing_available():
		return None
	filters = {}
	if cashier:
		filters["cashier"] = cashier
	if pos_profile:
		filters["pos_profile"] = pos_profile
	rows = frappe.get_all(
		REGISTER_CLOSING_DOCTYPE,
		filters=filters,
		fields=["name", "cashier", "pos_profile", "period_start", "period_end", "counted_cash", "creation"],
		order_by="period_end desc, creation desc",
		limit_page_length=1,
		ignore_permissions=True,
	)
	return rows[0] if rows else None


def _fp_order_datetime(order):
	try:
		return get_datetime(order.get("created_at"))
	except Exception:
		return None


def _fp_register_window_totals(start_dt, end_dt, cashier=None):
	start_date = str(getdate(start_dt))
	end_date = str(getdate(end_dt))
	orders = _management_collect_orders(
		date_from=start_date, date_to=end_date, source="all", cashier=cashier or None
	)

	start_dt = get_datetime(start_dt)
	end_dt = get_datetime(end_dt)
	method_totals = defaultdict(float)
	orders_count = 0
	total_sales = 0.0
	for order in orders:
		if not _is_revenue_order(order):
			continue
		order_dt = _fp_order_datetime(order)
		if order_dt and not (start_dt <= order_dt <= end_dt):
			continue
		method = (order.get("payment_method") or "").strip().lower()
		if method not in {"cash", "card", "credit"}:
			method = "other"
		amount = flt(order.get("grand_total"))
		method_totals[method] += amount
		total_sales += amount
		orders_count += 1

	return {
		"orders": orders_count,
		"total_sales": flt(total_sales),
		"cash_sales": flt(method_totals.get("cash")),
		"card_sales": flt(method_totals.get("card")),
		"credit_sales": flt(method_totals.get("credit")),
		"other_sales": flt(method_totals.get("other")),
	}


def _fp_register_window_for_user(cashier, pos_profile=None, date_from=None):
	last_closing = _fp_last_register_closing(cashier=cashier, pos_profile=pos_profile or None)
	if (
		pos_profile
		and not last_closing
	):
		# A brand-new register has no closing of its own yet; fall back to the
		# cashier's latest closing on any register so windows stay continuous.
		last_closing = _fp_last_register_closing(cashier=cashier)
	if last_closing and last_closing.get("period_end"):
		period_start = get_datetime(last_closing.get("period_end"))
	elif date_from:
		period_start = get_datetime(f"{getdate(date_from)} 00:00:00")
	else:
		period_start = get_datetime(f"{today()} 00:00:00")
	period_end = now_datetime()
	return last_closing, period_start, period_end


@frappe.whitelist()
def get_management_register_closing_summary(pos_profile=None, date_from=None):
	"""Live totals for the current open register window (before closing)."""
	_ensure_management_access()
	if not _fp_register_closing_available():
		frappe.throw(_("Restaurant Register Closing doctype is not installed yet. Run bench migrate."))

	cashier = frappe.session.user
	pos_profile = (pos_profile or "").strip()
	last_closing, period_start, period_end = _fp_register_window_for_user(
		cashier, pos_profile=pos_profile, date_from=date_from
	)
	totals = _fp_register_window_totals(period_start, period_end, cashier=cashier)

	shift_settings = _pos_shift_settings()
	suggested_float = flt(
		(last_closing.get("counted_cash") if last_closing else None)
		or shift_settings.get("opening", {}).get("cash_float")
		or 0
	)
	expected_cash = flt(suggested_float + totals.get("cash_sales"))

	return {
		"cashier": cashier,
		"pos_profile": pos_profile,
		"period_start": _json_safe_datetime(period_start),
		"period_end": _json_safe_datetime(period_end),
		"last_closing": last_closing or None,
		"totals": totals,
		"suggested_opening_float": suggested_float,
		"expected_cash": expected_cash,
		"tolerance": flt(shift_settings.get("closing", {}).get("tolerance") or 0),
		"currency": _get_currency(),
	}


def _fp_register_closing_receipt_html(closing_doc, breakdown):
	brand = _management_get_print_brand_settings()
	font_family = html_escape(str(brand.get("print_font_family") or "Peyda"))
	font_size = min(max(cint(brand.get("print_font_size") or 11), 8), 24)
	brand_name = html_escape(str(brand.get("brand_name") or ""))
	header_bg = html_escape(str(brand.get("header_background_color") or "#4A2522"))
	header_text = html_escape(str(brand.get("header_text_color") or "#FFFFFF"))
	currency = html_escape(_get_currency())

	def row(label, value):
		return (
			f"<tr><td>{html_escape(str(label))}</td>"
			f"<td class='num'>{html_escape(str(value))}</td></tr>"
		)

	rows_html = "".join(
		[
			row(_("Opening Float"), f"{flt(closing_doc.get('opening_float')):,.0f} {currency}"),
			row(_("Cash Sales"), f"{flt(closing_doc.get('cash_sales')):,.0f} {currency}"),
			row(_("Card Sales"), f"{flt(closing_doc.get('card_sales')):,.0f} {currency}"),
			row(_("Credit Sales"), f"{flt(closing_doc.get('credit_sales')):,.0f} {currency}"),
			row(_("Other Sales"), f"{flt(closing_doc.get('other_sales')):,.0f} {currency}"),
			row(_("Total Sales"), f"{flt(closing_doc.get('total_sales')):,.0f} {currency}"),
			row(_("Total Orders"), cint(closing_doc.get("total_orders"))),
			row(_("Expected Cash"), f"{flt(closing_doc.get('expected_cash')):,.0f} {currency}"),
			row(_("Counted Cash"), f"{flt(closing_doc.get('counted_cash')):,.0f} {currency}"),
			row(_("Difference"), f"{flt(closing_doc.get('cash_difference')):,.0f} {currency}"),
		]
	)

	return f"""
<div class="register-closing-receipt" dir="rtl" style="font-family:{font_family}, Peyda, Tahoma, sans-serif;font-size:{font_size}px;color:#1f3b32;">
  <div style="background:{header_bg};color:{header_text};padding:8px 12px;border-radius:10px 10px 0 0;text-align:center;">
    <strong>{brand_name}</strong><br />
    <small>{html_escape(str(_("Register Closing Receipt")))}</small>
  </div>
  <div style="border:1px dashed #c8bfb2;border-top:none;border-radius:0 0 10px 10px;padding:10px 12px;">
    <p style="margin:0 0 6px;">{html_escape(str(_("Closing No.")))}: {html_escape(str(closing_doc.get("name") or ""))}</p>
    <p style="margin:0 0 6px;">{html_escape(str(_("Cashier")))}: {html_escape(str(closing_doc.get("cashier") or ""))}</p>
    <p style="margin:0 0 8px;">{html_escape(str(_("Period")))}: {html_escape(str(closing_doc.get("period_start") or ""))} ← {html_escape(str(closing_doc.get("period_end") or ""))}</p>
    <table style="width:100%;border-collapse:collapse;" class="closing-table">
      <tbody>{rows_html}</tbody>
    </table>
  </div>
  <style>
    .closing-table td {{ border-bottom:1px dashed #e2d9cc; padding:4px 2px; }}
    .closing-table td.num {{ text-align:left; font-weight:600; }}
  </style>
</div>
"""


@frappe.whitelist()
def close_management_register(payload=None):
	"""Settle and close the cash register for the current cashier window."""
	_ensure_management_access()
	if not _fp_register_closing_available():
		frappe.throw(_("Restaurant Register Closing doctype is not installed yet. Run bench migrate."))

	data = _parse_json(payload, {})
	if not isinstance(data, dict):
		frappe.throw(_("Invalid payload format."))

	cashier = frappe.session.user
	pos_profile = (data.get("pos_profile") or "").strip()
	last_closing, period_start, period_end = _fp_register_window_for_user(
		cashier, pos_profile=pos_profile, date_from=data.get("date_from")
	)

	totals = _fp_register_window_totals(period_start, period_end, cashier=cashier)

	summary = get_management_register_closing_summary(pos_profile=pos_profile)
	opening_float = (
		flt(data.get("opening_float")) if data.get("opening_float") not in (None, "") else flt(summary.get("suggested_opening_float"))
	)
	counted_cash = flt(data.get("counted_cash") or 0)
	expected_cash = flt(opening_float + totals.get("cash_sales"))
	cash_difference = flt(counted_cash - expected_cash)
	tolerance = flt(summary.get("tolerance") or 0)
	within_tolerance = abs(cash_difference) <= tolerance if tolerance > 0 else cash_difference == 0

	doc = frappe.get_doc(
		{
			"doctype": REGISTER_CLOSING_DOCTYPE,
			"cashier": cashier,
			"pos_profile": pos_profile,
			"period_start": period_start,
			"period_end": period_end,
			"opening_float": opening_float,
			"total_sales": totals.get("total_sales"),
			"total_orders": totals.get("orders"),
			"cash_sales": totals.get("cash_sales"),
			"card_sales": totals.get("card_sales"),
			"credit_sales": totals.get("credit_sales"),
			"other_sales": totals.get("other_sales"),
			"expected_cash": expected_cash,
			"counted_cash": counted_cash,
			"cash_difference": cash_difference,
			"payments_breakdown_json": json.dumps(totals, ensure_ascii=False),
			"note": (data.get("note") or "").strip(),
		}
	).insert(ignore_permissions=True)
	frappe.db.commit()

	return {
		"status": "success",
		"closing": _fp_serialize_register_closing(doc),
		"within_tolerance": within_tolerance,
		"tolerance": tolerance,
		"receipt_html": _fp_register_closing_receipt_html(doc, totals),
	}


def _fp_serialize_register_closing(doc):
	return {
		"name": doc.get("name"),
		"cashier": doc.get("cashier") or "",
		"pos_profile": doc.get("pos_profile") or "",
		"period_start": _json_safe_datetime(doc.get("period_start")),
		"period_end": _json_safe_datetime(doc.get("period_end")),
		"opening_float": flt(doc.get("opening_float")),
		"total_sales": flt(doc.get("total_sales")),
		"total_orders": cint(doc.get("total_orders")),
		"cash_sales": flt(doc.get("cash_sales")),
		"card_sales": flt(doc.get("card_sales")),
		"credit_sales": flt(doc.get("credit_sales")),
		"other_sales": flt(doc.get("other_sales")),
		"expected_cash": flt(doc.get("expected_cash")),
		"counted_cash": flt(doc.get("counted_cash")),
		"cash_difference": flt(doc.get("cash_difference")),
		"note": doc.get("note") or "",
		"creation": _json_safe_datetime(doc.get("creation")),
	}


@frappe.whitelist()
def list_management_register_closings(limit=20, cashier=None):
	_ensure_management_access()
	if not _fp_register_closing_available():
		return {"closings": []}
	filters = {}
	if cashier:
		filters["cashier"] = cashier
	rows = frappe.get_all(
		REGISTER_CLOSING_DOCTYPE,
		filters=filters,
		fields=["name"],
		order_by="period_end desc, creation desc",
		limit_page_length=max(min(cint(limit) or 20, 100), 1),
		ignore_permissions=True,
	)
	closings = [_fp_serialize_register_closing(frappe.get_doc(REGISTER_CLOSING_DOCTYPE, row.get("name"))) for row in rows]
	return {"closings": closings, "currency": _get_currency()}


@frappe.whitelist()
def get_management_register_closing_detail(name=None):
	_ensure_management_access()
	if not _fp_register_closing_available():
		frappe.throw(_("Restaurant Register Closing doctype is not installed yet. Run bench migrate."))
	name = (name or "").strip()
	if not name or not frappe.db.exists(REGISTER_CLOSING_DOCTYPE, name):
		frappe.throw(_("Register closing not found."), frappe.DoesNotExistError)
	doc = frappe.get_doc(REGISTER_CLOSING_DOCTYPE, name)
	breakdown = _parse_json(doc.get("payments_breakdown_json") or "{}", {})
	return {
		"closing": _fp_serialize_register_closing(doc),
		"payments_breakdown": breakdown if isinstance(breakdown, dict) else {},
		"receipt_html": _fp_register_closing_receipt_html(doc, breakdown if isinstance(breakdown, dict) else {}),
	}


# ---------------------------------------------------------------------------
# Extra sales reports
# ---------------------------------------------------------------------------


def _fp_revenue_orders(orders):
	return [order for order in (orders or []) if _is_revenue_order(order)]


def _fp_item_category_map(item_codes):
	codes = sorted({code for code in (item_codes or []) if code})
	if not codes:
		return {}
	fields = ["name", "item_group"]
	has_restaurant_category = _has_column("Item", "restaurant_category")
	if has_restaurant_category:
		fields.append("restaurant_category")
	rows = frappe.get_all(
		"Item",
		filters={"name": ["in", codes]},
		fields=fields,
		ignore_permissions=True,
		limit_page_length=len(codes) + 5,
	)
	category_map = {}
	for row in rows:
		category = (row.get("restaurant_category") or "") if has_restaurant_category else ""
		category_map[row.get("name")] = category or (row.get("item_group") or "")
	return category_map


def _fp_item_group_title_map(group_names):
	names = sorted({name for name in (group_names or []) if name})
	if not names:
		return {}
	rows = frappe.get_all(
		"Item Group",
		filters={"name": ["in", names]},
		fields=["name", "item_group_name"],
		ignore_permissions=True,
		limit_page_length=len(names) + 5,
	)
	return {row.get("name"): (row.get("item_group_name") or row.get("name")) for row in rows}


def _fp_build_category_sales_rows(orders):
	revenue = _fp_revenue_orders(orders)
	item_codes = set()
	for order in revenue:
		for item in order.get("items") or []:
			item_code = (item.get("item_code") or "").strip()
			if item_code:
				item_codes.add(item_code)

	category_map = _fp_item_category_map(item_codes)
	title_map = _fp_item_group_title_map(category_map.values())

	grouped = defaultdict(lambda: {"category": "", "qty": 0.0, "sales": 0.0, "_order_keys": set()})
	total_sales = 0.0
	for order in revenue:
		for item in order.get("items") or []:
			item_code = (item.get("item_code") or "").strip()
			category_name = category_map.get(item_code) or ""
			category_title = title_map.get(category_name) or category_name or _("Uncategorized")
			bucket = grouped[category_title]
			bucket["category"] = category_title
			bucket["qty"] += flt(item.get("qty"))
			line_total = flt(item.get("line_total"))
			bucket["sales"] += line_total
			bucket["_order_keys"].add(order.get("name"))
			total_sales += line_total

	rows = []
	for bucket in grouped.values():
		rows.append(
			{
				"category": bucket["category"],
				"qty": round(bucket["qty"], 2),
				"sales": flt(bucket["sales"]),
				"orders": len(bucket["_order_keys"]),
				"share_percent": round(_safe_div(bucket["sales"] * 100, total_sales), 1),
			}
		)
	return sorted(rows, key=lambda row: row["sales"], reverse=True)


def _fp_order_table_label(order):
	table_label = (order.get("table_label") or "").strip()
	if table_label:
		return table_label
	table_name = (order.get("table") or "").strip()
	if table_name:
		return table_name
	customer_name = (order.get("customer_name") or "").strip()
	if (order.get("source") or "") == "table" and customer_name:
		return customer_name
	return ""


def _fp_build_table_sales_rows(orders):
	revenue = _fp_revenue_orders(orders)
	grouped = defaultdict(lambda: {"table": "", "orders": 0, "sales": 0.0, "items": 0.0})
	for order in revenue:
		table_label = _fp_order_table_label(order)
		if not table_label:
			continue
		bucket = grouped[table_label]
		bucket["table"] = table_label
		bucket["orders"] += 1
		bucket["sales"] += flt(order.get("grand_total"))
		bucket["items"] += sum(flt(item.get("qty")) for item in order.get("items") or [])
	return sorted(grouped.values(), key=lambda row: row["sales"], reverse=True)


def _fp_normalize_report_payment_method(order):
	method = (order.get("payment_method") or "").strip().lower()
	if method in {"cash", "card", "credit"}:
		return method
	if method:
		return "other"
	return "unknown"


def _fp_build_payment_method_rows(orders):
	revenue = _fp_revenue_orders(orders)
	grouped = defaultdict(lambda: {"method": "", "label": "", "orders": 0, "sales": 0.0})
	total_sales = 0.0
	for order in revenue:
		method = _fp_normalize_report_payment_method(order)
		bucket = grouped[method]
		bucket["method"] = method
		bucket["label"] = PAYMENT_METHOD_LABELS.get(method) or method
		bucket["orders"] += 1
		amount = flt(order.get("grand_total"))
		bucket["sales"] += amount
		total_sales += amount

	rows = []
	for bucket in grouped.values():
		rows.append(
			{
				**{key: value for key, value in bucket.items() if not key.startswith("_")},
				"share_percent": round(_safe_div(bucket["sales"] * 100, total_sales), 1),
			}
		)
	return sorted(rows, key=lambda row: row["sales"], reverse=True)


def _fp_build_product_sales_rows(orders):
	revenue = _fp_revenue_orders(orders)
	grouped = defaultdict(
		lambda: {"product_title": "", "qty": 0.0, "sales": 0.0, "_order_keys": set()}
	)
	total_sales = 0.0
	for order in revenue:
		for item in order.get("items") or []:
			title = (item.get("title") or item.get("item_code") or "").strip() or _("Untitled")
			bucket = grouped[title]
			bucket["product_title"] = title
			bucket["qty"] += flt(item.get("qty"))
			line_total = flt(item.get("line_total"))
			bucket["sales"] += line_total
			bucket["_order_keys"].add(order.get("name"))
			total_sales += line_total

	rows = []
	for bucket in grouped.values():
		qty = round(bucket["qty"], 2)
		rows.append(
			{
				"product_title": bucket["product_title"],
				"qty": qty,
				"sales": flt(bucket["sales"]),
				"orders": len(bucket["_order_keys"]),
				"avg_price": round(_safe_div(bucket["sales"], bucket["qty"]), 2),
				"share_percent": round(_safe_div(bucket["sales"] * 100, total_sales), 1),
			}
		)
	return sorted(rows, key=lambda row: row["sales"], reverse=True)


def _fp_build_shift_sales_rows(orders):
	shifts = fp_get_work_shifts()
	revenue = _fp_revenue_orders(orders)
	grouped = {}
	for shift in shifts:
		grouped[shift["key"]] = {
			"shift": shift["label"],
			"start": shift["start"],
			"end": shift["end"],
			"orders": 0,
			"sales": 0.0,
			"items": 0.0,
		}
	outside = {"shift": _("Outside Shifts"), "start": "", "end": "", "orders": 0, "sales": 0.0, "items": 0.0}

	for order in revenue:
		order_dt = _fp_order_datetime(order)
		shift = _fp_shift_for_datetime(shifts, order_dt) if order_dt else None
		bucket = grouped[shift["key"]] if shift else outside
		bucket["orders"] += 1
		bucket["sales"] += flt(order.get("grand_total"))
		bucket["items"] += sum(flt(item.get("qty")) for item in order.get("items") or [])

	rows = [grouped[shift["key"]] for shift in shifts]
	if outside["orders"] or outside["sales"]:
		rows.append(outside)
	return rows


@frappe.whitelist()
def get_management_report_category_sales(date_from=None, date_to=None):
	_ensure_management_access()
	orders = _management_collect_orders(date_from=date_from, date_to=date_to, source="all")
	rows = _fp_build_category_sales_rows(orders)
	summary = {
		"categories": len(rows),
		"total_sales": sum(flt(row["sales"]) for row in rows),
		"total_qty": round(sum(flt(row["qty"]) for row in rows), 2),
	}
	return _compose_management_report(
		"category-sales", "Category Sales", summary, rows,
		date_from=date_from, date_to=date_to, source="all", orders=orders,
	)


@frappe.whitelist()
def get_management_report_table_sales(date_from=None, date_to=None):
	_ensure_management_access()
	orders = _management_collect_orders(date_from=date_from, date_to=date_to, source="all")
	rows = _fp_build_table_sales_rows(orders)
	summary = {
		"tables": len(rows),
		"total_sales": sum(flt(row["sales"]) for row in rows),
		"total_orders": sum(cint(row["orders"]) for row in rows),
	}
	return _compose_management_report(
		"table-sales", "Table Sales", summary, rows,
		date_from=date_from, date_to=date_to, source="all", orders=orders,
	)


@frappe.whitelist()
def get_management_report_payment_methods(date_from=None, date_to=None):
	_ensure_management_access()
	orders = _management_collect_orders(date_from=date_from, date_to=date_to, source="all")
	rows = _fp_build_payment_method_rows(orders)
	summary = {
		"payment_methods": len(rows),
		"total_sales": sum(flt(row["sales"]) for row in rows),
		"total_orders": sum(cint(row["orders"]) for row in rows),
	}
	return _compose_management_report(
		"payment-methods", "Payment Methods", summary, rows,
		date_from=date_from, date_to=date_to, source="all", orders=orders,
	)


@frappe.whitelist()
def get_management_report_product_sales(date_from=None, date_to=None):
	_ensure_management_access()
	orders = _management_collect_orders(date_from=date_from, date_to=date_to, source="all")
	rows = _fp_build_product_sales_rows(orders)
	summary = {
		"products": len(rows),
		"total_sales": sum(flt(row["sales"]) for row in rows),
		"total_qty": round(sum(flt(row["qty"]) for row in rows), 2),
	}
	return _compose_management_report(
		"product-sales", "Product Sales", summary, rows,
		date_from=date_from, date_to=date_to, source="all", orders=orders,
	)


@frappe.whitelist()
def get_management_report_shift_sales(date_from=None, date_to=None):
	_ensure_management_access()
	_fp_ensure_ops_setting_fields()
	orders = _management_collect_orders(date_from=date_from, date_to=date_to, source="all")
	rows = _fp_build_shift_sales_rows(orders)
	summary = {
		"shifts": len([row for row in rows if row.get("orders") or row.get("sales")]),
		"total_sales": sum(flt(row["sales"]) for row in rows),
		"total_orders": sum(cint(row["orders"]) for row in rows),
	}
	return _compose_management_report(
		"shift-sales", "Shift Sales", summary, rows,
		date_from=date_from, date_to=date_to, source="all", orders=orders,
	)


def fp_build_report_bi(report_key, title, summary, rows, orders, previous_orders, meta):
	"""BI payload (KPIs, charts, tables, insights) for feature-pack reports."""
	kpis = []
	charts = []
	tables = []
	insights = []

	if report_key == "category-sales":
		prev_rows = _fp_build_category_sales_rows(previous_orders)
		prev_total = sum(flt(row.get("sales")) for row in prev_rows)
		total_sales = sum(flt(row.get("sales")) for row in rows)
		top_row = rows[0] if rows else {}
		kpis = [
			_bi_kpi("categories", _("Categories"), len(rows), "count", len(prev_rows)),
			_bi_kpi("total_sales", _("Total Sales"), total_sales, "money", prev_total),
			_bi_kpi(
				"top_category_sales",
				_("Top Category Sales"),
				flt(top_row.get("sales")) if top_row else 0,
				"money",
				0,
			),
		]
		labels = [row.get("category") for row in rows[:10]]
		charts = [
			{
				"key": "category-sales",
				"title": _("Sales by Category"),
				"type": "bar",
				"unit": "money",
				"labels": labels,
				"series": [
					{
						"key": "sales",
						"label": _("Sales"),
						"color": "#2f6f5c",
						"values": [flt(row.get("sales")) for row in rows[:10]],
					}
				],
			},
			{
				"key": "category-share",
				"title": _("Category Share %"),
				"type": "bar",
				"unit": "percent",
				"labels": labels,
				"series": [
					{
						"key": "share",
						"label": _("Share %"),
						"color": "#3e8ed0",
						"values": [flt(row.get("share_percent")) for row in rows[:10]],
					}
				],
			},
		]
		tables = [
			{"key": "category-table", "title": _("Category Sales"), "columns": _table_columns_from_rows(rows), "rows": rows}
		]
		if top_row:
			insights.append(
				{
					"key": "top-category",
					"severity": "info",
					"text": _("Best selling category: {0} ({1}% of sales).").format(
						top_row.get("category"), top_row.get("share_percent")
					),
				}
			)

	elif report_key == "table-sales":
		prev_rows = _fp_build_table_sales_rows(previous_orders)
		prev_total = sum(flt(row.get("sales")) for row in prev_rows)
		total_sales = sum(flt(row.get("sales")) for row in rows)
		top_row = rows[0] if rows else {}
		kpis = [
			_bi_kpi("tables", _("Active Tables"), len(rows), "count", len(prev_rows)),
			_bi_kpi("total_sales", _("Table Sales"), total_sales, "money", prev_total),
			_bi_kpi(
				"top_table_sales",
				_("Top Table Sales"),
				flt(top_row.get("sales")) if top_row else 0,
				"money",
				0,
			),
		]
		labels = [row.get("table") for row in rows[:10]]
		charts = [
			{
				"key": "table-sales",
				"title": _("Sales by Table"),
				"type": "bar",
				"unit": "money",
				"labels": labels,
				"series": [
					{
						"key": "sales",
						"label": _("Sales"),
						"color": "#2f6f5c",
						"values": [flt(row.get("sales")) for row in rows[:10]],
					}
				],
			}
		]
		tables = [
			{"key": "table-sales-table", "title": _("Table Sales"), "columns": _table_columns_from_rows(rows), "rows": rows}
		]
		if not rows:
			insights.append(
				{
					"key": "no-table-sales",
					"severity": "info",
					"text": _("No table (dine-in) sales were found in this window."),
				}
			)

	elif report_key == "payment-methods":
		prev_rows = _fp_build_payment_method_rows(previous_orders)
		prev_total = sum(flt(row.get("sales")) for row in prev_rows)
		total_sales = sum(flt(row.get("sales")) for row in rows)
		cash_row = next((row for row in rows if row.get("method") == "cash"), {})
		credit_row = next((row for row in rows if row.get("method") == "credit"), {})
		kpis = [
			_bi_kpi("total_sales", _("Total Sales"), total_sales, "money", prev_total),
			_bi_kpi("cash_sales", _("Cash Sales"), flt(cash_row.get("sales")), "money", 0),
			_bi_kpi("credit_sales", _("Credit Sales"), flt(credit_row.get("sales")), "money", 0),
			_bi_kpi("methods", _("Payment Methods"), len(rows), "count", len(prev_rows)),
		]
		labels = [row.get("label") for row in rows]
		charts = [
			{
				"key": "payment-method-sales",
				"title": _("Sales by Payment Method"),
				"type": "bar",
				"unit": "money",
				"labels": labels,
				"series": [
					{
						"key": "sales",
						"label": _("Sales"),
						"color": "#2f6f5c",
						"values": [flt(row.get("sales")) for row in rows],
					}
				],
			},
			{
				"key": "payment-method-orders",
				"title": _("Orders by Payment Method"),
				"type": "bar",
				"unit": "count",
				"labels": labels,
				"series": [
					{
						"key": "orders",
						"label": _("Orders"),
						"color": "#c08a2a",
						"values": [cint(row.get("orders")) for row in rows],
					}
				],
			},
		]
		tables = [
			{
				"key": "payment-method-table",
				"title": _("Payment Method Breakdown"),
				"columns": _table_columns_from_rows(rows),
				"rows": rows,
			}
		]
		unknown_row = next((row for row in rows if row.get("method") == "unknown"), None)
		if unknown_row and unknown_row.get("sales"):
			insights.append(
				{
					"key": "unknown-methods",
					"severity": "warn",
					"text": _(
						"Part of the sales has no recorded payment method (for example table orders paid outside the POS)."
					),
				}
			)

	elif report_key == "product-sales":
		prev_rows = _fp_build_product_sales_rows(previous_orders)
		prev_total = sum(flt(row.get("sales")) for row in prev_rows)
		total_sales = sum(flt(row.get("sales")) for row in rows)
		top_row = rows[0] if rows else {}
		bottom_row = rows[-1] if rows else {}
		kpis = [
			_bi_kpi("products", _("Products Sold"), len(rows), "count", len(prev_rows)),
			_bi_kpi("total_sales", _("Total Sales"), total_sales, "money", prev_total),
			_bi_kpi(
				"top_product_sales",
				_("Top Product Sales"),
				flt(top_row.get("sales")) if top_row else 0,
				"money",
				0,
			),
		]
		labels = [row.get("product_title") for row in rows[:10]]
		charts = [
			{
				"key": "product-sales-top",
				"title": _("Top Products"),
				"type": "bar",
				"unit": "money",
				"labels": labels,
				"series": [
					{
						"key": "sales",
						"label": _("Sales"),
						"color": "#2f6f5c",
						"values": [flt(row.get("sales")) for row in rows[:10]],
					}
				],
			}
		]
		if len(rows) > 10:
			bottom_labels = [row.get("product_title") for row in rows[-10:]]
			charts.append(
				{
					"key": "product-sales-bottom",
					"title": _("Least Selling Products"),
					"type": "bar",
					"unit": "money",
					"labels": bottom_labels,
					"series": [
						{
							"key": "sales",
							"label": _("Sales"),
							"color": "#b84f4f",
							"values": [flt(row.get("sales")) for row in rows[-10:]],
						}
					],
				}
			)
		tables = [
			{"key": "product-sales-table", "title": _("Product Sales Detail"), "columns": _table_columns_from_rows(rows), "rows": rows}
		]
		if top_row:
			insights.append(
				{
					"key": "top-product",
					"severity": "info",
					"text": _("Best selling product: {0} ({1}% of sales).").format(
						top_row.get("product_title"), top_row.get("share_percent")
					),
				}
			)
		if bottom_row and len(rows) > 1:
			insights.append(
				{
					"key": "bottom-product",
					"severity": "warn",
					"text": _("Least selling product: {0}.").format(bottom_row.get("product_title")),
				}
			)

	elif report_key == "shift-sales":
		prev_rows = _fp_build_shift_sales_rows(previous_orders)
		prev_total = sum(flt(row.get("sales")) for row in prev_rows)
		total_sales = sum(flt(row.get("sales")) for row in rows)
		best_shift = max(rows, key=lambda row: row.get("sales") or 0) if rows else {}
		kpis = [
			_bi_kpi("total_sales", _("Total Sales"), total_sales, "money", prev_total),
			_bi_kpi(
				"total_orders",
				_("Total Orders"),
				sum(cint(row.get("orders")) for row in rows),
				"count",
				sum(cint(row.get("orders")) for row in prev_rows),
			),
			_bi_kpi(
				"best_shift_sales",
				_("Best Shift Sales"),
				flt(best_shift.get("sales")) if best_shift else 0,
				"money",
				0,
			),
		]
		labels = [row.get("shift") for row in rows]
		charts = [
			{
				"key": "shift-sales",
				"title": _("Sales by Work Shift"),
				"type": "bar",
				"unit": "money",
				"labels": labels,
				"series": [
					{
						"key": "sales",
						"label": _("Sales"),
						"color": "#2f6f5c",
						"values": [flt(row.get("sales")) for row in rows],
					}
				],
			},
			{
				"key": "shift-orders",
				"title": _("Orders by Work Shift"),
				"type": "bar",
				"unit": "count",
				"labels": labels,
				"series": [
					{
						"key": "orders",
						"label": _("Orders"),
						"color": "#3e8ed0",
						"values": [cint(row.get("orders")) for row in rows],
					}
				],
			},
		]
		tables = [
			{"key": "shift-sales-table", "title": _("Shift Sales"), "columns": _table_columns_from_rows(rows), "rows": rows}
		]
		if best_shift and best_shift.get("sales"):
			insights.append(
				{
					"key": "best-shift",
					"severity": "info",
					"text": _("Best shift: {0} ({1}-{2}).").format(
						best_shift.get("shift"), best_shift.get("start"), best_shift.get("end")
					),
				}
			)

	return {"kpis": kpis, "charts": charts, "tables": tables, "insights": insights}


# ---------------------------------------------------------------------------
# Printer font settings
# ---------------------------------------------------------------------------


def _fp_print_font_payload():
	settings = _management_get_print_brand_settings()
	return {
		"font_family": settings.get("print_font_family") or "Peyda",
		"font_size": min(max(cint(settings.get("print_font_size") or 11), 8), 24),
		"receipt_font_scale": settings.get("print_receipt_font_scale") or "متوسط",
		"font_options": list(PRINT_FONT_FAMILY_OPTIONS),
		"scale_options": list(PRINT_RECEIPT_SCALE_OPTIONS),
	}


@frappe.whitelist()
def get_management_print_font_settings():
	_ensure_management_access()
	return _fp_print_font_payload()


@frappe.whitelist()
def set_management_print_font_settings(payload=None):
	_ensure_management_access()
	data = _parse_json(payload, {})
	if not isinstance(data, dict):
		frappe.throw(_("Invalid payload format."))

	font_family = (data.get("font_family") or "").strip()
	if font_family:
		if font_family not in PRINT_FONT_FAMILY_OPTIONS:
			frappe.throw(
				_("Unsupported font: {0}. Allowed: {1}").format(
					font_family, ", ".join(PRINT_FONT_FAMILY_OPTIONS)
				)
			)
		frappe.db.set_single_value(
			"Restaurant Print Brand Settings", "print_font_family", font_family, update_modified=True
		)

	if data.get("font_size") not in (None, ""):
		font_size = min(max(cint(data.get("font_size")), 8), 24)
		frappe.db.set_single_value(
			"Restaurant Print Brand Settings", "print_font_size", font_size, update_modified=True
		)

	receipt_scale = (data.get("receipt_font_scale") or "").strip()
	if receipt_scale:
		if receipt_scale not in PRINT_RECEIPT_SCALE_OPTIONS:
			frappe.throw(
				_("Unsupported receipt scale: {0}.").format(receipt_scale)
			)
		frappe.db.set_single_value(
			"Restaurant Print Brand Settings", "print_receipt_font_scale", receipt_scale, update_modified=True
		)

	frappe.db.commit()
	return {"status": "success", "settings": _fp_print_font_payload()}


# ---------------------------------------------------------------------------
# Dashboard layout persistence
# ---------------------------------------------------------------------------


def _fp_normalize_dashboard_layout(layout):
	layout = layout if isinstance(layout, dict) else {}
	widgets_raw = layout.get("widgets") if isinstance(layout.get("widgets"), dict) else layout
	widgets = {}
	for index, key in enumerate(DASHBOARD_WIDGET_KEYS):
		entry = widgets_raw.get(key) if isinstance(widgets_raw.get(key), dict) else {}
		widgets[key] = {
			"visible": bool(entry.get("visible", True)),
			"order": cint(entry.get("order", index)),
		}
	return {"widgets": widgets}


@frappe.whitelist()
def get_management_dashboard_layout():
	_ensure_management_access()
	try:
		raw = frappe.defaults.get_user_default("restaurant_dashboard_layout") or ""
	except Exception:
		raw = ""
	layout = _parse_json(raw, {})
	return {"layout": _fp_normalize_dashboard_layout(layout)}


@frappe.whitelist()
def set_management_dashboard_layout(payload=None):
	_ensure_management_access()
	data = _parse_json(payload, {})
	layout_input = data.get("layout") if isinstance(data, dict) and "layout" in data else data
	normalized = _fp_normalize_dashboard_layout(layout_input)
	frappe.defaults.set_user_default(
		"restaurant_dashboard_layout",
		json.dumps(normalized, ensure_ascii=False),
		user=frappe.session.user,
	)
	frappe.db.commit()
	return {"status": "success", "layout": normalized}


def _fp_register_into_api_module():
	"""Expose this module's public API on ``restaurant.api``.

	``api.py`` re-exports everything via ``from restaurant.api_feature_pack import *``
	so the endpoints resolve as ``/api/method/restaurant.api.<name>`` — but if this
	module happens to be imported *before* ``api.py`` (e.g. by the v2_7 patch during
	``bench migrate``), that star-import runs against a partially-initialized module
	and silently copies nothing. Pushing the names here, after ``__all__`` and all
	defs exist, keeps ``restaurant.api.<endpoint>`` resolvable in both import orders.
	"""
	import sys

	api_module = sys.modules.get("restaurant.api")
	if api_module is None:
		return
	for _name in __all__:
		if _name in globals():
			setattr(api_module, _name, globals()[_name])


_fp_register_into_api_module()
